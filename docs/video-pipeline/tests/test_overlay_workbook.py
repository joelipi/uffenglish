"""Unit tests for overlay_workbook.py (stdlib only; openpyxl optional).

Covers the MAIN + INPUT planner: column classification, the step-key rule,
formula shapes, subtitle concatenation round-trips against the real seeded
master, the wider production header shape, and an end-to-end .xlsx write
(skipped when openpyxl is unavailable).
"""

import csv
import os
import sys
import tempfile
import unittest

# Make the sibling module importable when this file runs standalone or under
# discovery from the repo root (mirrors test_pipeline_lib.py / run_all.py).
HERE = os.path.dirname(os.path.abspath(__file__))
PIPELINE_DIR = os.path.dirname(HERE)
sys.path.insert(0, PIPELINE_DIR)

try:
    import openpyxl  # noqa: F401
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from overlay_workbook import (
    build_input_tables,
    build_workbook_plan,
    classify_columns,
    excel_col_letter,
    group_member_map,
    input_sections,
    layout_input_tables,
    read_csv_rows,
    sibling_groups,
    split_phrases_out,
    step_key_for_row,
    subtitle_concat_formula,
    subtitle_groups,
    subtitle_is_reconstructable,
    vlookup_formula,
)

SEED_CSV = os.path.join(PIPELINE_DIR, "master-seeded.csv")


def tiny_rows():
    headers = ["title_text", "order", "character", "phrase", "filename", "join",
               "video_file", "subtitle_text", "course_id", "lesson_id", "lesson_title",
               "response_type", "next_step", "choose_step_text", "srt_es"]

    def row(order, phrase, filename, video_file, lesson, lesson_title, response, subtitle):
        return {"title_text": "T", "order": order, "character": "C", "phrase": phrase,
                "filename": filename, "join": "", "video_file": video_file, "course_id": "c",
                "lesson_id": lesson, "lesson_title": lesson_title, "response_type": response,
                "next_step": "", "choose_step_text": "", "srt_es": "", "subtitle_text": subtitle}

    block = [
        # Step v (lesson a): two takes whose overlay embeds take 1, then takes 1-2.
        row("1", "P one", "v01", "v", "a", "Ask", "r",
            "<aside><strong><b>A1</b> P one</strong><br><b>A2</b> ___________</aside>"),
        row("2", "P two", "v02", "v", "a", "Ask", "r",
            "<aside><b>A1</b> P one<br><strong><b>A2</b> P two</strong></aside>"),
        # Step w (same lesson a, other response_type): single take, stored overlay.
        row("1", "P three", "w01", "w", "a", "Ask", "r2",
            "<aside><strong><b>A1</b> P three.</strong></aside>"),
        # Step z (lesson b): single take, stored overlay.
        row("1", "P four", "z01", "z", "b", "Answer", "r",
            "<aside><strong><b>A1</b> P four.</strong></aside>"),
    ]
    return headers, block


class StepKeyTest(unittest.TestCase):
    def test_join_wins_over_video_file(self):
        self.assertEqual(step_key_for_row({"join": "j", "video_file": "v"}), "j")

    def test_falls_back_to_video_file(self):
        self.assertEqual(step_key_for_row({"join": "", "video_file": "v"}), "v")

    def test_blank_both_is_blank(self):
        self.assertEqual(step_key_for_row({"join": "", "video_file": ""}), "")

    def test_whitespace_is_content_matching_the_sheet_formula(self):
        # The MAIN `_step` formula uses an exact `=""` blank test, so 2 spaces
        # are content (step key "  "), not blank — Python must agree.
        self.assertEqual(step_key_for_row({"join": "  ", "video_file": "v"}), "  ")


class EmptyAndDuplicateInputTest(unittest.TestCase):
    def test_empty_rows_raise_cleanly(self):
        headers = ["filename", "phrase", "video_file", "join", "lesson_id"]
        with self.assertRaises(ValueError):
            build_workbook_plan(headers, [])

    def test_duplicate_filenames_raise(self):
        headers, rows = tiny_rows()
        dupes = list(rows) + [dict(rows[0], phrase="P dup")]
        with self.assertRaises(ValueError) as ctx:
            build_workbook_plan(headers, dupes)
        self.assertIn("duplicate", str(ctx.exception))
        self.assertIn("filename", str(ctx.exception))

    def test_blank_filenames_are_not_duplicates(self):
        headers, rows = tiny_rows()
        blank = [dict(r, filename="") for r in rows]
        # Blank keys carry no VLOOKUP identity; they must not trip the guard.
        plan = build_workbook_plan(headers, blank)
        self.assertEqual(len(plan["main"]), len(blank))

    def test_key_materializes_as_formula(self):
        headers, rows = tiny_rows()
        plan = build_workbook_plan(headers, rows)
        header = plan["main_header"]
        for col in ("lesson_id", "video_file", "join"):
            kind, payload = plan["main"][0][header.index(col)]
            self.assertEqual(kind, "formula", col)
            self.assertTrue(payload.startswith("=IFERROR(VLOOKUP("), col)


class ColumnLettersTest(unittest.TestCase):
    def test_letters(self):
        self.assertEqual(excel_col_letter(1), "A")
        self.assertEqual(excel_col_letter(26), "Z")
        self.assertEqual(excel_col_letter(27), "AA")
        self.assertEqual(excel_col_letter(70), "BR")


class ClassifyTest(unittest.TestCase):
    def test_tiny_sheet_kinds(self):
        headers, rows = tiny_rows()
        kinds = classify_columns(headers, rows)
        self.assertEqual(kinds["filename"], "value")
        self.assertEqual(kinds["order"], "value")
        self.assertEqual(kinds["phrase"], "phrase")
        self.assertEqual(kinds["subtitle_text"], "subtitle")
        self.assertEqual(kinds["title_text"], "global")
        self.assertEqual(kinds["lesson_title"], "lesson")
        self.assertEqual(kinds["lesson_id"], "key")
        self.assertEqual(kinds["video_file"], "key")
        # Constant inside the one step -> step kind.
        self.assertEqual(kinds["response_type"], "step")
        # Blank file-wide -> blank kind (MAIN stays blank, no formula).
        self.assertEqual(kinds["next_step"], "blank")

    def test_production_header_shape_classifies(self):
        headers, rows = tiny_rows()
        headers = headers + ["choose_step_next", "choose_step_text_es", "publish_lesson_id",
                             "srt_pt", "srt_fr", "srt_bn", "phrase_es"]
        for row in rows:
            for extra in ("choose_step_next", "choose_step_text_es", "publish_lesson_id",
                          "srt_pt", "srt_fr", "srt_bn", "phrase_es"):
                row[extra] = ""
        kinds = classify_columns(headers, rows)
        self.assertEqual(kinds["phrase_es"], "phrase")
        plan = build_workbook_plan(headers, rows)
        self.assertIn("_step", plan["main_header"])
        self.assertEqual(len(plan["main"]), len(rows))

    def test_seeded_master_kinds(self):
        headers, rows = read_csv_rows(SEED_CSV)
        kinds = classify_columns(headers, rows)
        for col in ("title_text", "character", "course_id", "success_srt"):
            self.assertEqual(kinds[col], "global", col)
        for col in ("lesson_title", "recap_sources", "intro_video"):
            self.assertEqual(kinds[col], "lesson", col)
        for col in ("join", "video_file", "lesson_id"):
            self.assertEqual(kinds[col], "key", col)
        self.assertEqual(kinds["phrase"], "phrase")
        self.assertEqual(kinds["subtitle_text"], "subtitle")
        # Only the take-level columns vary inside a step.
        self.assertEqual(sorted(k for k, v in kinds.items() if v == "per_row"), [])


class FormulaShapeTest(unittest.TestCase):
    def test_vlookup_is_exact_match_with_fallback(self):
        formula = vlookup_formula("$P2", "INPUT!$A$10:$I$118", 3)
        self.assertTrue(formula.startswith("="))
        self.assertIn("VLOOKUP($P2,INPUT!$A$10:$I$118,3,FALSE)", formula)
        self.assertIn("IFERROR", formula)

    def test_step_helper_uses_exact_blank_test_not_trim(self):
        # TRIM would collapse internal double spaces and diverge from the CSV
        # rule, so the `_step` formula must test `=""` exactly.
        headers, rows = tiny_rows()
        plan = build_workbook_plan(headers, rows)
        kind, payload = plan["main"][0][plan["main_header"].index("_step")]
        self.assertEqual(kind, "formula")
        self.assertNotIn("TRIM", payload)
        self.assertRegex(payload, r'^=IF\(\$[A-Z]+\$2="",\$[A-Z]+\$2,\$[A-Z]+\$2\)$')


class SubtitleSplitTest(unittest.TestCase):
    def test_split_and_rejoin_round_trip(self):
        subtitle = "<aside><b>A1</b> P one<br><strong><b>A2</b> P two</strong></aside>"
        ok, fragments = split_phrases_out(subtitle, ["P one", "P two"])
        self.assertTrue(ok)
        self.assertEqual(fragments[0] + "P one" + fragments[1] + "P two" + fragments[2], subtitle)

    def test_missing_phrase_fails(self):
        ok, _fragments = split_phrases_out("<aside>x</aside>", ["P one"])
        self.assertFalse(ok)

    def test_concat_formula_references_phrase_cells(self):
        subtitle = "<aside><strong><b>A1</b> P one</strong><br><b>A2</b> ___________</aside>"
        formula = subtitle_concat_formula(subtitle, [("P one", "$H$2")])
        self.assertTrue(formula.startswith("="))
        self.assertIn("$H$2", formula)
        self.assertIn("___________", formula)
        self.assertNotIn("P one\"&", formula.replace("$H$2", ""))

    def test_tiny_block_is_reconstructable(self):
        headers, rows = tiny_rows()
        groups = subtitle_groups(rows)
        ok, ordered = subtitle_is_reconstructable(0, rows, groups)
        self.assertTrue(ok)
        self.assertEqual(ordered, [0])
        ok, ordered = subtitle_is_reconstructable(1, rows, groups)
        self.assertTrue(ok)
        self.assertEqual(ordered, [0, 1])

    def test_joined_halves_are_separate_overlays(self):
        # The two halves of a join share a step_key but each video_file block
        # carries its own four-take overlay.
        _headers, rows = tiny_rows()
        halves = [dict(rows[0], filename="h01", video_file="h_i", join="h", phrase="Half one",
                       subtitle_text="<aside><strong>Half one</strong>___________</aside>"),
                  dict(rows[0], filename="h02", video_file="h_i", join="h", order="2",
                       phrase="Half two",
                       subtitle_text="<aside>Half one<strong>Half two</strong></aside>"),
                  dict(rows[0], filename="j01", video_file="h_ii", join="h", phrase="Other one",
                       subtitle_text="<aside><strong>Other one</strong>___________</aside>")]
        groups = subtitle_groups(halves)
        self.assertEqual(len(groups), 2)  # h_i x2 takes, h_ii x1 take
        ok, ordered = subtitle_is_reconstructable(0, halves, groups)
        self.assertTrue(ok)
        self.assertEqual(ordered, [0])
        ok, _ordered = subtitle_is_reconstructable(2, halves, groups)
        self.assertFalse(ok)  # single-take group -> stored overlay

    def test_single_row_step_is_not_reconstructable(self):
        _headers, rows = tiny_rows()
        solo = [dict(rows[0], filename="solo", video_file="solo", join="intro", phrase="")]
        solo[0]["subtitle_text"] = "<aside><strong><b>A1</b> P one.</strong></aside>"
        groups = subtitle_groups(solo)
        ok, _ordered = subtitle_is_reconstructable(0, solo, groups)
        self.assertFalse(ok)

    def test_blank_video_file_rows_stand_alone(self):
        _headers, rows = tiny_rows()
        intros = [dict(rows[0], filename="i1", video_file="", join="intro", phrase="",
                       subtitle_text="<aside><strong>Intro one.</strong></aside>"),
                  dict(rows[0], filename="i2", video_file="", join="intro", phrase="",
                       subtitle_text="<aside><strong>Intro two.</strong></aside>")]
        groups = subtitle_groups(intros)
        self.assertEqual(len(groups), 2)  # blank video_file rows never merge
        for index in (0, 1):
            ok, _ordered = subtitle_is_reconstructable(index, intros, groups)
            self.assertFalse(ok)

    def test_intro_rows_fall_back_to_stored_overlays(self):
        headers, rows = tiny_rows()
        intros = [dict(rows[0], filename="i1", video_file="", join="intro", phrase="",
                       subtitle_text="<aside><strong>Intro one.</strong></aside>"),
                  dict(rows[0], filename="i2", video_file="", join="intro", phrase="",
                       subtitle_text="<aside><strong>Intro two.</strong></aside>")]
        plan = build_workbook_plan(headers, intros)
        self.assertEqual(len(plan["tables"]["subtitles_rows"]), 2)
        header = plan["main_header"]
        kind, payload = plan["main"][0][header.index("subtitle_text")]
        self.assertEqual(kind, "formula")
        self.assertIn("VLOOKUP", payload)
        self.assertNotIn("&", payload)
        self.assertEqual(plan["expected"][0][header.index("subtitle_text")],
                         "<aside><strong>Intro one.</strong></aside>")


class StoreOnceInvariantTest(unittest.TestCase):
    def test_no_input_literal_is_embedded_in_a_main_formula(self):
        # The whole point: a value that lives on INPUT appears on MAIN only as
        # a reference, never as a re-typed literal.
        headers, rows = read_csv_rows(SEED_CSV)
        plan = build_workbook_plan(headers, rows)
        header = plan["main_header"]
        global_cols = [c for c, k in plan["kinds"].items() if k == "global"]
        self.assertTrue(global_cols)
        for rix, _row in enumerate(rows):
            for col in global_cols:
                kind, payload = plan["main"][rix][header.index(col)]
                self.assertEqual(kind, "formula")
                value = rows[rix][col]
                if value:
                    # The literal text never appears inside the formula.
                    self.assertNotIn(value, payload, (rix, col))

    def test_shared_text_is_written_to_input_once(self):
        headers, rows = read_csv_rows(SEED_CSV)
        plan = build_workbook_plan(headers, rows)
        globals_values = [value for _col, value in plan["tables"]["globals"]]
        self.assertEqual(len(globals_values), len(set(globals_values)))
        # Each lesson appears once in the lessons table, not per row.
        lesson_ids = [r[0] for r in plan["tables"]["lessons_rows"]]
        self.assertEqual(len(lesson_ids), len(set(lesson_ids)))
        step_keys = plan["tables"]["step_keys"]
        self.assertEqual(len(step_keys), len(set(step_keys)))


class LayoutSingleSourceTest(unittest.TestCase):
    def test_sections_carry_the_planned_rows(self):
        headers, rows = read_csv_rows(SEED_CSV)
        plan = build_workbook_plan(headers, rows)
        layout = plan["layout"]
        for name, section in layout["sections"].items():
            self.assertEqual(section["header_row"], section["label_row"] + 1, name)
            self.assertEqual(section["first_data"], section["header_row"] + 1, name)
            self.assertEqual(section["last_data"],
                             section["first_data"] + len(section["rows"]) - 1, name)
            self.assertEqual(layout["%s_range" % name], section["range"], name)

    def test_layout_is_recomputable_from_tables_alone(self):
        headers, rows = read_csv_rows(SEED_CSV)
        plan = build_workbook_plan(headers, rows)
        self.assertEqual(layout_input_tables(plan["tables"]), plan["layout"])
        names = [s["name"] for s in input_sections(plan["tables"])]
        self.assertEqual(names[:4], ["globals", "lessons", "steps", "phrases"])


class MemberMapTest(unittest.TestCase):
    def test_member_map_round_trips_every_row(self):
        _headers, rows = tiny_rows()
        groups = subtitle_groups(rows)
        member_of = group_member_map(groups)
        self.assertEqual(sorted(member_of), list(range(len(rows))))
        for index, members in member_of.items():
            self.assertIn(index, members)


class PlanFidelityTest(unittest.TestCase):
    def test_expected_values_reproduce_the_csv(self):
        headers, rows = tiny_rows()
        plan = build_workbook_plan(headers, rows)
        for rix, row in enumerate(rows):
            for cix, col in enumerate(plan["main_header"]):
                if col == "_step":
                    self.assertEqual(plan["expected"][rix][cix], step_key_for_row(row))
                else:
                    self.assertEqual(plan["expected"][rix][cix], row[col], (rix, col))

    def test_content_cells_are_formulas_not_copies(self):
        headers, rows = tiny_rows()
        plan = build_workbook_plan(headers, rows)
        header = plan["main_header"]
        for rix in range(len(rows)):
            for col in ("title_text", "character", "phrase", "subtitle_text",
                        "lesson_title", "response_type", "video_file", "lesson_id"):
                kind, payload = plan["main"][rix][header.index(col)]
                self.assertEqual(kind, "formula", col)
                self.assertTrue(payload.startswith("="), col)
        # Row keys stay plain values.
        kind, payload = plan["main"][0][header.index("filename")]
        self.assertEqual((kind, payload), ("value", "v01"))

    def test_seeded_master_plan_is_faithful(self):
        headers, rows = read_csv_rows(SEED_CSV)
        plan = build_workbook_plan(headers, rows)
        tables = plan["tables"]
        # Deduplication: one lesson row per lesson, one step per join/video_file.
        self.assertEqual(len(tables["lessons_rows"]), 6)
        self.assertEqual(len(tables["steps_rows"]), len(tables["step_keys"]))
        self.assertEqual(len(tables["phrases_rows"]), len(rows))
        # Every expected cell reproduces the CSV.
        for rix, row in enumerate(rows):
            for cix, col in enumerate(plan["main_header"]):
                if col == "_step":
                    continue
                self.assertEqual(plan["expected"][rix][cix], row[col], (rix, col))
        # Subtitle formulas evaluate back to the CSV overlay: substitute each
        # phrase reference with its expected phrase value and compare.
        header = plan["main_header"]
        phrase_col = header.index("phrase")
        for rix, row in enumerate(rows):
            kind, payload = plan["main"][rix][header.index("subtitle_text")]
            if kind != "formula" or "&" not in payload:
                continue  # stored-overlay fallback row
            evaluated = payload[1:]  # strip leading '='
            # Replace each sibling phrase ref with the quoted phrase text.
            take_groups = tables["take_groups"]
            ok, ordered = subtitle_is_reconstructable(rix, rows, take_groups)
            self.assertTrue(ok)
            for member in ordered:
                ref = "$%s$%d" % (excel_col_letter(phrase_col + 1), member + 2)
                phrase = rows[member]["phrase"].replace('"', '""')
                evaluated = evaluated.replace(ref, '"%s"' % phrase)
            # What remains is a chain of quoted literals joined by &.
            parts = [p for p in evaluated.split("&")]
            self.assertTrue(parts)
            joined = "".join(p[1:-1].replace('""', '"') for p in parts)
            self.assertEqual(joined, row["subtitle_text"], rix)


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl not installed")
class WriteWorkbookTest(unittest.TestCase):
    def test_write_and_read_back(self):
        from overlay_workbook import write_workbook
        headers, rows = tiny_rows()
        plan = build_workbook_plan(headers, rows)
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "workbook.xlsx")
            write_workbook(plan, out)
            book = openpyxl.load_workbook(out)
            self.assertIn("MAIN", book.sheetnames)
            self.assertIn("INPUT", book.sheetnames)
            main = book["MAIN"]
            self.assertEqual(main["A1"].value, headers[0])
            self.assertEqual(main.cell(row=1, column=len(headers) + 1).value, "_step")
            # A repeated global is a formula pointing at INPUT, not a copy.
            title_col = headers.index("title_text") + 1
            self.assertTrue(str(main.cell(row=2, column=title_col).value).startswith("=INPUT!"))
            # Row keys are plain values.
            filename_col = headers.index("filename") + 1
            self.assertEqual(main.cell(row=2, column=filename_col).value, "v01")
            # INPUT carries the single global value once.
            ws_in = book["INPUT"]
            found = [c.value for row in ws_in.iter_rows(min_row=1, max_col=2) for c in row]
            self.assertIn("title_text", found)
            self.assertIn("T", found)

    def test_seeded_workbook_write(self):
        from overlay_workbook import write_workbook
        headers, rows = read_csv_rows(SEED_CSV)
        plan = build_workbook_plan(headers, rows)
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "seeded.xlsx")
            write_workbook(plan, out)
            book = openpyxl.load_workbook(out)
            main = book["MAIN"]
            self.assertEqual(main.max_row, len(rows) + 1)
            self.assertEqual(main.max_column, len(headers) + 1)

    def test_planned_ranges_match_written_tables(self):
        import re
        from overlay_workbook import write_workbook
        headers, rows = read_csv_rows(SEED_CSV)
        plan = build_workbook_plan(headers, rows)
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "ranges.xlsx")
            write_workbook(plan, out)
            book = openpyxl.load_workbook(out)
            ws_in = book["INPUT"]
            for name, first_header_cell in (("lessons", "lesson_id"), ("steps", "step_key"),
                                            ("phrases", "filename")):
                rng = plan["layout"]["%s_range" % name]
                match = re.fullmatch(r"INPUT!\$A\$(\d+):\$[A-Z]+\$(\d+)", rng)
                self.assertIsNotNone(match, rng)
                header_row, last_row = int(match.group(1)), int(match.group(2))
                self.assertEqual(ws_in.cell(row=header_row, column=1).value,
                                 first_header_cell, name)
                table = plan["tables"]["%s_rows" % name]
                self.assertEqual(last_row, header_row + len(table), name)
                # Global value cells land on the planned addresses too.
                for col, addr in plan["layout"]["global_cells"].items():
                    gmatch = re.fullmatch(r"\$B\$(\d+)", addr)
                    self.assertIsNotNone(gmatch, addr)
                    self.assertEqual(ws_in.cell(row=int(gmatch.group(1)), column=1).value, col)


if __name__ == "__main__":
    unittest.main()
